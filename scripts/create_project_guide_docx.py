from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "HUONG_DAN_DU_AN.txt"
OUTPUT = ROOT / "HUONG_DAN_DU_AN_SALES_LAB.docx"

FALLBACK_TEXT = """HƯỚNG DẪN DỰ ÁN SALES LAB ĐA NGÀNH
==================================

Phiên bản tài liệu: 22/09/2026
Website: https://taki-sales-lab.crisnguyen3090.chatgpt.site/

1. MỤC TIÊU DỰ ÁN
-----------------

Sales Lab Đa Ngành là hệ thống luyện kỹ năng tư vấn bán hàng cho đội ngũ sale. Người học đóng vai sale; hệ thống đóng vai khách hàng khó tính, phản hồi theo ngành hàng, chân dung khách, tình huống và câu trả lời gần nhất của sale.

Hệ thống giúp sale hỏi đúng nhu cầu, phản hồi đúng điều khách vừa nói, xử lý phản đối và chốt bước tiếp theo. Quản lý có thể theo dõi điểm, lịch sử hội thoại và lỗi của từng nhân sự.

2. ĐƯỜNG DẪN VÀ TÀI KHOẢN
-------------------------

Website: https://taki-sales-lab.crisnguyen3090.chatgpt.site/

Người dùng đăng nhập bằng tài khoản ChatGPT. Lần đầu sử dụng, nhập họ tên và đội hoặc phòng ban. Hệ thống không tự tạo hoặc lưu mật khẩu riêng.

Quyền admin được xác định bằng email trong biến môi trường TAKI_ADMIN_EMAIL. Admin xem được các phiên luyện gần nhất của toàn bộ người dùng, gồm điểm, chỉ số chi tiết, transcript và phần chữa bài.

3. CÁC NGÀNH ĐANG HỖ TRỢ
------------------------

1. Thời trang.
2. Phụ kiện.
3. Mỹ phẩm và chăm sóc cá nhân.
4. Thực phẩm và đồ uống.
5. Đồ gia dụng.
6. Điện tử và công nghệ.
7. Giáo dục và khóa học.
8. Du lịch và lưu trú.
9. Bất động sản.
10. Y tế và chăm sóc sức khỏe.
11. Spa và thẩm mỹ.
12. Nội thất.
13. Vật liệu xây dựng.

Mỗi ngành có sản phẩm hoặc dịch vụ mẫu cùng nhiều chân dung khách. Mỗi chân dung có hoàn cảnh, nhu cầu, nỗi sợ, ngân sách, thời điểm mua, trải nghiệm cũ và người ra quyết định riêng.

4. TÌNH HUỐNG LUYỆN TẬP
-----------------------

- Khách lạnh, thiếu kiên nhẫn.
- Khách bận, muốn quyết nhanh.
- Khách phản đối giá.
- Khách mất niềm tin.
- Khách nghi ngờ mức độ phù hợp.
- Khách chưa phải người quyết định cuối.
- Khách đang so sánh đối thủ.
- Khách hiểu ngành và hỏi sâu.

5. CÁCH SALE SỬ DỤNG
--------------------

1. Đăng nhập và hoàn tất hồ sơ cá nhân.
2. Chọn ngành hàng.
3. Chọn sản phẩm hoặc dịch vụ.
4. Chọn tình huống muốn luyện.
5. Chọn chân dung khách hàng.
6. Bấm Bắt đầu luyện.
7. Trả lời khách như trong cuộc chat bán hàng thật.
8. Đọc phần Chữa câu trả lời sau mỗi lượt.
9. Bấm Kết thúc và lưu điểm.
10. Mở Lịch sử của tôi để xem lại phiên đã lưu.

Phần chữa bài gồm điểm câu trả lời, lỗi cần sửa, câu trả lời gợi ý và bước tương ứng trong quy trình TAKI.

6. NGUYÊN TẮC PHẢN HỒI CỦA KHÁCH HÀNG
--------------------------------------

Khách phải phản hồi câu sale vừa nói trước khi hỏi vấn đề mới. Khi sale hỏi đúng, khách cung cấp dữ kiện thuộc chân dung đã chọn rồi mới phản biện tiếp.

Nếu sale hỏi nhiều ý, khách có thể trả lời tối đa ba dữ kiện liên quan trong cùng một lượt. Khách không được tự chuyển sang chủ đề không xuất hiện trong lời sale.

Nếu sale hỏi lệch ngành, khách hỏi lại lý do của câu hỏi và kéo cuộc trò chuyện về đúng nhu cầu. Điểm câu trả lời bị giới hạn ở mức rất thấp.

Ví dụ: khách mua trang sức nhưng sale hỏi chiều cao hoặc cân nặng. Khách phải phản ứng rằng thông tin đó liên quan thế nào tới việc chọn trang sức, thay vì tự chuyển sang hỏi bảo hành hay giao hàng.

Nếu sale đưa lựa chọn mơ hồ như mẫu 1 có được không nhưng chưa cung cấp thông tin, khách phải yêu cầu làm rõ mẫu nào, đặc điểm gì và vì sao phù hợp.

Nếu sale trả lời vô nghĩa, khách yêu cầu trả lời lại đúng trọng tâm và câu đó gần như không được tính điểm.

7. NGUYÊN TẮC CHẤM ĐIỂM
-----------------------

Điểm từng câu dựa trên các yếu tố sau:

- Trả lời đúng điều khách vừa nói.
- Thể hiện sự lắng nghe hoặc đồng cảm.
- Có câu hỏi khai thác cụ thể.
- Dùng dữ kiện đúng ngành.
- Đưa bằng chứng hoặc số liệu khi cần.
- Chốt bước tiếp theo đúng giai đoạn.
- Không quá dài, chung chung hoặc lệch chủ đề.

Điểm tổng phiên là trung bình điểm các câu sale. Phiên luyện lưu bốn chỉ số: bám sát lời khách, đúng quy trình, câu hỏi khai thác và xử lý tự nhiên.

8. QUY TRÌNH BÁN HÀNG TAKI
--------------------------

1. Mở đầu và xin phép.
2. Tạo thiện cảm, kiểm soát cuộc trò chuyện.
3. Khai thác hiện trạng và mục tiêu.
4. Đào sâu điểm đau và hậu quả.
5. Nối giải pháp đúng nhu cầu.
6. Trình bày bằng chứng phù hợp.
7. Xử lý từ chối theo CETAA.
8. Chốt bước tiếp theo cụ thể.

9. HƯỚNG DẪN CHO ADMIN VÀ QUẢN LÝ
---------------------------------

Sau khi đăng nhập bằng email admin, mở tab Quản trị. Admin có thể xem người luyện, đội nhóm, ngành, sản phẩm, tình huống, điểm và số lượt chat; đồng thời mở transcript để đọc từng câu và phần chữa bài.

Đề xuất vận hành:

- Mỗi sale luyện ít nhất ba phiên cho tình huống quan trọng.
- Không đánh giá chỉ bằng điểm tổng; cần đọc các câu bị chấm thấp.
- Luyện lại cùng tình huống với chân dung khách khác để tránh học thuộc.
- Dùng transcript trong buổi coaching một một hoặc họp đội.

10. CẬP NHẬT NGÀNH SẢN PHẨM VÀ INSIGHT
--------------------------------------

Dữ liệu chính nằm trong app/data.ts. Khi thêm hoặc sửa ngành, cần có tên ngành, bối cảnh, tiêu chí quyết định, bằng chứng khách yêu cầu, câu hỏi chẩn đoán, phản đối về giá, niềm tin, hậu mãi, so sánh đối thủ, từ khóa đặc thù và ngân hàng câu hỏi khó.

Mỗi chân dung khách cần có hoàn cảnh, nhu cầu cốt lõi, nỗi sợ, ngân sách, thời điểm mua, trải nghiệm trước đây, người ra quyết định và phong cách nói.

Mỗi sản phẩm phải gắn đúng industryId để xuất hiện trong đúng ngành.

11. KIỂM THỬ BẮT BUỘC
---------------------

Mỗi ngành phải được kiểm tra với ít nhất các nhóm sau:

1. Sale hỏi nhu cầu hoặc mục tiêu.
2. Sale hỏi ngân sách hoặc giá.
3. Sale hỏi thời gian mua.
4. Sale hỏi trải nghiệm trước đây.
5. Sale hỏi người ra quyết định.
6. Sale nhắc bằng chứng hoặc chứng nhận.
7. Sale nhắc bảo hành hoặc hậu mãi.
8. Sale so sánh đối thủ.
9. Sale hỏi câu thuộc ngành khác.
10. Sale trả lời vô nghĩa hoặc quá ngắn.
11. Sale đưa lựa chọn mơ hồ như mẫu 1 có được không.
12. Sale thúc chốt khi chưa đủ dữ kiện.

Tiêu chí đạt:

- Khách trả lời đúng câu sale vừa hỏi.
- Câu mơ hồ bị yêu cầu làm rõ.
- Câu lệch ngành bị phản ứng trực tiếp và chấm điểm thấp.
- Không lặp nguyên văn câu hỏi cũ.
- Không sinh dữ kiện trái chân dung.
- Phần chữa bài sửa đúng lỗi vừa xảy ra.

12. CÔNG NGHỆ VÀ FILE QUAN TRỌNG
-------------------------------

- React và TypeScript.
- Next.js và Vinext.
- Cloudflare Workers.
- Cloudflare D1 và Drizzle ORM.
- ChatGPT Sign In.
- OpenAI Sites.

Các file chính:

- app/data.ts: ngành, sản phẩm, tình huống và insight khách.
- app/training-portal.tsx: luồng chat, phản hồi và chấm điểm.
- app/api/me/route.ts: hồ sơ người dùng.
- app/api/sessions/route.ts: lịch sử cá nhân.
- app/api/admin/sessions/route.ts: dashboard admin.
- db/schema.ts: cấu trúc cơ sở dữ liệu.
- drizzle/: migration cơ sở dữ liệu.
- .openai/hosting.json: cấu hình Sites và D1.

13. CHẠY VÀ KIỂM TRA TRÊN MÁY
-----------------------------

Yêu cầu Node.js phiên bản 22.13.0 trở lên.

npm ci
npm run build
npm run dev

Khi chạy cục bộ với D1, áp dụng migration trong thư mục drizzle trước khi kiểm tra đăng ký, lưu lịch sử và dashboard admin.

14. LƯU Ý KỸ THUẬT
------------------

Phiên bản hiện tại dùng bộ máy hội thoại theo quy tắc và dữ liệu insight trong mã nguồn. Hệ thống chưa gọi mô hình ngôn ngữ bên ngoài ở mỗi lượt chat. Vì vậy ngành mới phải có từ khóa, insight, câu hỏi khó và kiểm thử hồi quy đầy đủ.

Không đưa mật khẩu, token, khóa API hoặc dữ liệu khách hàng thật vào mã nguồn. Với sức khỏe, bất động sản, tài chính và thẩm mỹ, nội dung mô phỏng không được cam kết chắc chắn và phải yêu cầu bằng chứng hoặc điều kiện cụ thể.

15. QUY TRÌNH XUẤT BẢN
----------------------

1. Cập nhật mã nguồn và dữ liệu.
2. Chạy kiểm thử hội thoại.
3. Chạy npm run build và sửa hết lỗi.
4. Commit đúng phiên bản đã kiểm tra.
5. Đóng gói kết quả build.
6. Lưu phiên bản mới trên OpenAI Sites.
7. Triển khai phiên bản đã lưu.
8. Kiểm tra trạng thái triển khai thành công.
9. Test lại tình huống quan trọng trên website thật.

16. CHECKLIST NGHIỆM THU
------------------------

[ ] Đăng nhập được.
[ ] Người mới đăng ký được hồ sơ.
[ ] Chọn được đủ 13 ngành.
[ ] Sản phẩm và chân dung thay đổi đúng theo ngành.
[ ] Bắt đầu và gửi hội thoại được.
[ ] Khách phản hồi trực tiếp câu sale vừa nói.
[ ] Câu lệch ngành bị phản ứng và chấm điểm thấp.
[ ] Câu mơ hồ bị yêu cầu làm rõ.
[ ] Phần chữa bài đúng lỗi.
[ ] Lưu được phiên luyện.
[ ] Người dùng xem được lịch sử cá nhân.
[ ] Admin xem được lịch sử và transcript.
[ ] Giao diện sử dụng được trên máy tính và điện thoại.
"""


def set_run_font(run, name="Arial", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)
        borders.append(tag)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=9, color=(90, 101, 117))
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1559B5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(r_pr)
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_text_with_url(paragraph, text):
    match = re.search(r"https?://\S+", text)
    if not match:
        run = paragraph.add_run(text)
        set_run_font(run, size=11)
        return
    before, url, after = text[:match.start()], match.group(0), text[match.end():]
    if before:
        run = paragraph.add_run(before)
        set_run_font(run, size=11)
    add_hyperlink(paragraph, url, url)
    if after:
        run = paragraph.add_run(after)
        set_run_font(run, size=11)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(31, 41, 55)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size in (("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True

    if "Checklist" not in [s.name for s in doc.styles]:
        checklist = doc.styles.add_style("Checklist", WD_STYLE_TYPE.PARAGRAPH)
        checklist.base_style = normal
        checklist.paragraph_format.left_indent = Inches(0.2)
        checklist.paragraph_format.first_line_indent = Inches(-0.2)
        checklist.paragraph_format.space_after = Pt(4)


def add_cover(doc):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(72)
    p.add_run("Hướng dẫn dự án Sales Lab Đa Ngành")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(14)
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("Tài liệu vận hành, cập nhật và kiểm thử hệ thống luyện sales")
    set_run_font(run, size=15, color=(45, 55, 72))

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(24)
    run = intro.add_run(
        "Tài liệu này giúp đội sale, quản lý và người phụ trách hệ thống sử dụng Sales Lab, "
        "đọc kết quả huấn luyện, cập nhật dữ liệu đa ngành và kiểm thử phản hồi khách hàng mô phỏng trước khi xuất bản."
    )
    set_run_font(run, size=12)

    table = doc.add_table(rows=3, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.55)
    table.columns[1].width = Inches(4.7)
    rows = [
        ("Phiên bản", "22/09/2026"),
        ("Phạm vi", "13 ngành, 26 sản phẩm dịch vụ mẫu và 26 chân dung khách"),
        ("Website", "https://taki-sales-lab.crisnguyen3090.chatgpt.site/"),
    ]
    for idx, (label, value) in enumerate(rows):
        left, right = table.rows[idx].cells
        for cell in (left, right):
            set_cell_margins(cell)
            set_cell_borders(cell)
        shade_cell(left, "E8F0FE")
        r = left.paragraphs[0].add_run(label)
        set_run_font(r, size=10.5, bold=True)
        if value.startswith("http"):
            add_hyperlink(right.paragraphs[0], value, value)
        else:
            r = right.paragraphs[0].add_run(value)
            set_run_font(r, size=10.5)

    doc.add_page_break()


def main():
    lines = (SOURCE.read_text(encoding="utf-8") if SOURCE.exists() else FALLBACK_TEXT).splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    configure_styles(doc)
    add_cover(doc)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("SALES LAB ĐA NGÀNH   |   HƯỚNG DẪN DỰ ÁN")
    set_run_font(run, size=8.5, bold=True, color=(90, 101, 117))
    add_page_number(section.footer.paragraphs[0])

    headings = []
    for i, line in enumerate(lines):
        if i + 1 < len(lines) and re.fullmatch(r"[-=]+", lines[i + 1].strip()):
            if re.match(r"^\d+\.\s", line):
                headings.append(line.strip())

    toc_title = doc.add_paragraph("Nội dung tài liệu", style="Heading 1")
    toc_title.paragraph_format.space_before = Pt(0)
    for heading in headings:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.32)
        p.paragraph_format.first_line_indent = Inches(-0.32)
        p.paragraph_format.space_after = Pt(2)
        number = re.match(r"^(\d+)\.", heading).group(1)
        text = re.sub(r"^\d+\.\s*", "", heading).title()
        run = p.add_run(f"{number}. {text}")
        set_run_font(run, size=10.5)
    doc.add_page_break()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        next_is_rule = i + 1 < len(lines) and bool(re.fullmatch(r"[-=]+", lines[i + 1].strip()))

        if i < 5:
            i += 1
            continue
        if not stripped or re.fullmatch(r"[-=]+", stripped):
            i += 1
            continue
        if next_is_rule:
            heading_text = re.sub(r"^\d+\.\s*", "", stripped).title()
            if heading_text.upper() != "KẾT THÚC TÀI LIỆU":
                doc.add_paragraph(heading_text, style="Heading 1")
            i += 2
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_text_with_url(p, stripped[2:])
        elif re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.28)
            add_text_with_url(p, stripped)
        elif stripped.startswith("[ ]"):
            p = doc.add_paragraph(style="Checklist")
            add_text_with_url(p, "☐ " + stripped[3:].strip())
        else:
            p = doc.add_paragraph()
            add_text_with_url(p, stripped)
        i += 1

    doc.core_properties.title = "Hướng dẫn dự án Sales Lab Đa Ngành"
    doc.core_properties.subject = "Hướng dẫn vận hành, cập nhật và kiểm thử hệ thống Sales Lab"
    doc.core_properties.author = "TAKI"
    doc.core_properties.keywords = "Sales Lab, TAKI, đào tạo sales, đa ngành"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
